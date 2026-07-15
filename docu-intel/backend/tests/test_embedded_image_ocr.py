from __future__ import annotations

import io
import zipfile
from pathlib import Path

from PIL import Image

from app.parsers.embedded_images import EmbeddedImage, extract_embedded_image_pages
from app.parsers.types import ExtractedBlock, ExtractedDocument, ExtractedPage


def _png_bytes() -> bytes:
    image = Image.new("RGB", (16, 16), "white")
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def test_embedded_image_ocr_persists_only_meaningful_text(tmp_path: Path, monkeypatch):
    def fake_parse_image(path: Path, output_dir: Path, ocr_engine):
        return ExtractedDocument(
            pages=[
                ExtractedPage(
                    page_number=1,
                    text="referencia visible",
                    blocks=[ExtractedBlock("text", "referencia visible", 1)],
                )
            ]
        )

    monkeypatch.setattr("app.parsers.embedded_images.parse_image", fake_parse_image)

    pages = extract_embedded_image_pages(
        [EmbeddedImage("photo.png", _png_bytes())],
        output_dir=tmp_path,
        ocr_engine=object(),
        first_page_number=2,
    )

    assert len(pages) == 1
    assert pages[0].page_number == 2
    assert pages[0].blocks[0].page_number == 2
    assert pages[0].text.startswith("[Imagen incrustada: photo.png]")
    assert list((tmp_path / "embedded").glob("*.png"))


def test_non_image_attachments_do_not_consume_the_image_limit(tmp_path: Path, monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "max_embedded_images_per_document", 1)
    monkeypatch.setattr(
        "app.parsers.embedded_images.parse_image",
        lambda path, output_dir, ocr_engine: ExtractedDocument(
            pages=[ExtractedPage(page_number=1, text="texto de imagen")]
        ),
    )

    pages = extract_embedded_image_pages(
        [
            EmbeddedImage("notes.txt", b"not an image"),
            EmbeddedImage("photo.png", _png_bytes()),
        ],
        output_dir=tmp_path,
        ocr_engine=object(),
        first_page_number=1,
    )

    assert len(pages) == 1
    assert "photo.png" in pages[0].text


def test_embedded_logo_is_searchable_but_not_an_ocr_quality_page(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(
        "app.parsers.embedded_images.parse_image",
        lambda path, output_dir, ocr_engine: ExtractedDocument(
            pages=[
                ExtractedPage(
                    page_number=1,
                    text="Decoraciones Egea",
                    ocr_confidence=0.50,
                    ocr_engine="vision",
                )
            ]
        ),
    )

    pages = extract_embedded_image_pages(
        [EmbeddedImage("image001.jpg", _png_bytes())],
        output_dir=tmp_path,
        ocr_engine=object(),
        first_page_number=1,
    )

    assert pages[0].ocr_content_kind == "decorative"
    assert pages[0].ocr_confidence is None
    assert "Decoraciones Egea" in pages[0].text


def test_docx_media_is_discovered_without_changing_the_document(tmp_path: Path):
    from app.parsers.docx import _embedded_images

    docx = tmp_path / "sample.docx"
    with zipfile.ZipFile(docx, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("word/media/image1.png", _png_bytes())

    images = _embedded_images(docx)

    assert [(image.filename, image.content) for image in images] == [("image1.png", _png_bytes())]


def test_excel_media_is_discovered(tmp_path: Path):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ExcelImage

    from app.parsers.excel import _embedded_images

    image_path = tmp_path / "image.png"
    image_path.write_bytes(_png_bytes())
    workbook = Workbook()
    workbook.active.add_image(ExcelImage(str(image_path)), "A1")
    spreadsheet = tmp_path / "sample.xlsx"
    workbook.save(spreadsheet)
    workbook.close()

    images = _embedded_images(spreadsheet)

    assert len(images) == 1
    assert images[0].filename.endswith(".png")
    assert images[0].content.startswith(b"\x89PNG")


def test_msg_attachment_collects_image_bytes_only():
    from app.parsers.msg import _image_attachment

    image_attachment = type("Attachment", (), {"longFilename": "photo.jpeg", "data": b"jpeg-bytes"})()
    non_binary_attachment = type("Attachment", (), {"longFilename": "notes.txt", "data": None})()

    assert _image_attachment(image_attachment) == EmbeddedImage("photo.jpeg", b"jpeg-bytes")
    assert _image_attachment(non_binary_attachment) is None
